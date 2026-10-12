"""Create a synthetic Word source for protected-input validation.

The generated content is deliberately fictional. It is rendered and inspected
before a separate Microsoft Word operation encrypts a copy for tests.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def create(output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    for style_name in ("Title", "Heading 1"):
        style = doc.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = "Aptos"
        paragraph_properties = style.element.pPr
        if paragraph_properties is not None:
            border = paragraph_properties.find(qn("w:pBdr"))
            if border is not None:
                paragraph_properties.remove(border)

    doc.add_paragraph("Synthetic Conveyor Recovery Procedure", style="Title")
    doc.add_paragraph(
        "This fictional source tests whether a protected Word file can be opened "
        "with authorization and whether its facts survive Wiki publication. "
        "The equipment values below are test values, not operating instructions."
    )

    doc.add_paragraph("Sequence and dependencies", style="Heading 1")
    doc.add_paragraph(
        "When controller C17 reports fault F09, isolate conveyor lane B before "
        "loading the last verified snapshot. After the restore, compare the "
        "sensor S4 reading with the PLC event stream. Resume movement only if "
        "both sources remain aligned for 90 seconds. Record the snapshot ID "
        "and the operator approval before restarting the lane."
    )

    doc.add_paragraph("Thresholds and exceptions", style="Heading 1")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Signal"
    table.rows[0].cells[1].text = "Test threshold"
    table.rows[0].cells[2].text = "Required action"
    for signal, threshold, action in (
        ("Motor current", "Above 14.2 A", "Keep lane B isolated"),
        ("Vibration", "Above 3.8 mm/s", "Inspect bearing housing"),
        ("Telemetry gap", "Over 5 seconds", "Repeat S4 to PLC comparison"),
    ):
        cells = table.add_row().cells
        cells[0].text = signal
        cells[1].text = threshold
        cells[2].text = action
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.add_paragraph("Evidence chain", style="Heading 1")
    doc.add_paragraph(
        "Sensor S4 feeds PLC P2, and PLC P2 authorizes conveyor lane B. "
        "A stale S4 sample must not be treated as evidence of a healthy lane. "
        "The source identifier for this test is WORD-CANARY-10521."
    )
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    create(args.output.resolve())


if __name__ == "__main__":
    main()
