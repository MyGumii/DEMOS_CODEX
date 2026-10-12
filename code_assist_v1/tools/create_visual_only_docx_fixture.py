"""Create an isolated Word fixture whose decisive facts exist only in a drawing.

This file is for 10521 source-fidelity regression tests, not a company DRM sample.
It refuses to overwrite an existing fixture.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUTPUT = Path(__file__).resolve().parents[1] / "output/docx/synthetic-visual-only-10521.docx"


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for path in (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/segoeui.ttf")):
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def main() -> None:
    if OUTPUT.exists():
        raise SystemExit(f"Refusing to replace the existing fixture: {OUTPUT}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    image = Image.new("RGB", (1600, 760), "white")
    draw = ImageDraw.Draw(image)
    title_font, body_font = _font(40), _font(29)
    draw.text((80, 52), "DOCX VISUAL CANARY 10521", fill="#172554", font=title_font)
    stages = (
        ("1  Detect", ("Telemetry absent", "more than 19 ms")),
        ("2  Recover", ("Verify snapshot digest", "before restore")),
        ("3  Release", ("Supervisor J03", "approves release")),
    )
    for index, (heading, detail_lines) in enumerate(stages):
        left = 70 + index * 510
        draw.rounded_rectangle((left, 180, left + 455, 600), radius=24,
                               outline="#1d4ed8", width=5, fill="#eff6ff")
        draw.text((left + 24, 225), heading, fill="#1e3a8a", font=title_font)
        for line_number, detail in enumerate(detail_lines):
            draw.text((left + 24, 340 + 48 * line_number), detail,
                      fill="#111827", font=body_font)
        if index < 2:
            right = left + 455
            draw.line((right + 5, 390, right + 44, 390), fill="#1d4ed8", width=7)
            draw.polygon(((right + 44, 390), (right + 27, 378), (right + 27, 402)),
                         fill="#1d4ed8")

    picture = BytesIO()
    image.save(picture, format="PNG")
    picture.seek(0)
    document = Document()
    document.sections[0].page_width = Inches(8.5)
    document.sections[0].page_height = Inches(11)
    title_style = document.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.size = Pt(20)
    style_properties = title_style._element.pPr
    if style_properties is not None:
        for border in style_properties.xpath("./w:pBdr"):
            style_properties.remove(border)
    title = document.add_paragraph("Controller recovery diagram", style="Title")
    document.add_paragraph(
        "The decision order, timing condition, and release approval are shown only in the diagram below."
    )
    document.add_picture(picture, width=Inches(6.5))
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
