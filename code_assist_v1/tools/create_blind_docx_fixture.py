"""Create a fictional Word source without Wiki-writing or test instructions.

The facts are intentionally spread across prose and two tables. This file is
only a source fixture; its expected answers are kept in the separate audit.
"""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output/docx/conveyor-controller-return-review-10521.docx"


def _table(document: Document, headers: tuple[str, ...], rows: tuple[tuple[str, ...], ...]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, heading in enumerate(headers):
        table.rows[0].cells[index].text = heading
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite a source fixture: {OUTPUT}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 17), ("Heading 1", 12), ("Heading 2", 10.5)):
        styles[name].font.name = "Aptos"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_properties = styles["Title"]._element.pPr
    title_border = title_properties.find(qn("w:pBdr")) if title_properties is not None else None
    if title_border is not None:
        title_properties.remove(title_border)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Conveyor Controller Return Review")
    document.add_paragraph(
        "This review records the 18 September controller recovery checks for lanes B through J. "
        "The release decision is lane-specific: normal readings on one lane do not authorize "
        "a neighboring lane. Only D02 and D08 received signed release decisions in this record."
    )

    document.add_paragraph("Decision basis", style="Heading 1")
    document.add_paragraph(
        "The incident record assigns each controller, lane and snapshot. Before release, the "
        "snapshot checksum must match the signed incident value and a valid timebase must cover "
        "the complete 90 s observation window. Motor current must be strictly below 14.2 A, "
        "bearing vibration strictly below 3.8 mm/s and every telemetry gap strictly below 5 s. "
        "Equality with a limit fails. A missing sample makes the window incomplete even when "
        "the largest gap reported by the received samples is under 5 s. These conditions are "
        "conjunctive; an operator note cannot substitute for a measured or signed condition."
    )

    document.add_paragraph("Measurements", style="Heading 1")
    document.add_paragraph(
        "All current and vibration numbers below are the worst observed values for the named "
        "90 s window. The gap column is computed from received telemetry; completeness is recorded "
        "separately in the evidence register."
    )
    _table(document, ("Case", "Controller and lane", "Snapshot", "Current A", "Vibration mm/s", "Gap s"), (
        ("D01", "C17 / B", "S42", "13.5", "3.9", "4"),
        ("D02", "C18 / C", "S43", "13.7", "3.4", "3"),
        ("D03", "C19 / D", "S44", "14.2", "3.2", "2"),
        ("D04", "C20 / E", "S45", "12.9", "3.1", "5"),
        ("D05", "C21 / F", "S46", "13.0", "3.0", "2"),
        ("D06", "C22 / G", "S47", "13.1", "3.2", "3"),
        ("D07", "C23 / H", "S48", "13.8", "3.7", "4"),
        ("D08", "C24 / J", "S49", "13.6", "3.2", "2"),
    ))

    document.add_paragraph("Evidence register and signed outcome", style="Heading 1")
    _table(document, ("Case", "Timebase", "Snapshot check", "Window", "Signed outcome"), (
        ("D01", "Valid", "Match", "Complete", "HOLD"),
        ("D02", "Valid", "Match", "Complete", "RELEASE"),
        ("D03", "Valid", "Match", "Complete", "HOLD"),
        ("D04", "Valid", "Match", "Complete", "HOLD"),
        ("D05", "Invalid", "Match", "Complete", "HOLD"),
        ("D06", "Valid", "Mismatch", "Complete", "HOLD"),
        ("D07", "Valid", "Match", "Incomplete", "HOLD"),
        ("D08", "Valid", "Match", "Complete", "RELEASE"),
    ))

    document.add_paragraph("Discrepancy notes", style="Heading 1")
    document.add_paragraph(
        "D01 exceeded the bearing vibration limit at 3.9 mm/s. A maintenance note later "
        "reported 3.6 mm/s after a bearing replacement at 14:20, but no new signed 90 s trace "
        "or release decision accompanied that note. D01 therefore remains HOLD in this review."
    )
    document.add_paragraph(
        "D03 reached 14.2 A exactly; the strict current comparison does not permit an equality "
        "exception. D04 reached a 5 s telemetry gap exactly. D05 had a clock reset during the "
        "window, making its timebase invalid despite three favorable numeric readings."
    )
    document.add_paragraph(
        "D06 used snapshot S47, but its calculated checksum ended 7B42 while the signed "
        "incident record ended 7B24. D07 lacks the sample expected at second 18; the four-second "
        "gap shown in the measurements table covers received samples only. Both remain on hold."
    )
    document.add_paragraph(
        "D02 met every condition for lane C, and D08 met every condition for lane J. These signed "
        "releases do not clear lanes B, D, E, F, G or H. Any later correction requires a complete "
        "new observation and its own authorized decision, rather than a revision of the original "
        "recorded outcome."
    )

    document.save(OUTPUT)
    print(f"source={OUTPUT}")
    print(f"sha256={sha256(OUTPUT.read_bytes()).hexdigest()}")


if __name__ == "__main__":
    main()
