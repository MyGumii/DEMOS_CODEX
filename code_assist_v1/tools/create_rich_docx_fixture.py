"""Create a new synthetic DOCX for 10521 upload and Wiki-fidelity tests.

All controllers, readings and decisions are fictional test data, not equipment
operating guidance. The output deliberately contains paragraphs and a real Word
table so extraction and authoring can be checked independently.
"""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "output" / "docx" / "synthetic-conveyor-rich-10521.docx"
CANARY = "RICH-DOCX-CANARY-10521"

CASES = (
    ("D01", "C17", "B", "S42", "13.5 A", "3.9 mm/s", "4 s", "HOLD", "bearing vibration exceeded 3.8 mm/s"),
    ("D02", "C18", "C", "S43", "13.7 A", "3.4 mm/s", "3 s", "RELEASE", "all three measurements were strictly below limits"),
    ("D03", "C19", "D", "S44", "14.2 A", "3.2 mm/s", "2 s", "HOLD", "motor current equalled its strict limit"),
    ("D04", "C20", "E", "S45", "12.9 A", "3.1 mm/s", "5 s", "HOLD", "the largest telemetry gap equalled its strict limit"),
    ("D05", "C21", "F", "S46", "13.0 A", "3.0 mm/s", "2 s", "HOLD", "the trace timebase was invalid"),
    ("D06", "C22", "G", "S47", "13.1 A", "3.2 mm/s", "3 s", "HOLD", "snapshot checksum did not match the incident record"),
)


def _cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def _cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for name in ("top", "left", "bottom", "right"):
        side = OxmlElement(f"w:{name}")
        side.set(qn("w:val"), "single")
        side.set(qn("w:sz"), "4")
        side.set(qn("w:color"), "D9D9D9")
        borders.append(side)
    tc_pr.append(borders)


def _cell_padding(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, twips in (("top", 90), ("bottom", 90), ("left", 110), ("right", 110)):
        side = OxmlElement(f"w:{name}")
        side.set(qn("w:w"), str(twips))
        side.set(qn("w:type"), "dxa")
        margins.append(side)
    tc_pr.append(margins)


def _table(doc: Document) -> None:
    headings = ("Case", "Controller", "Lane", "Snapshot", "Current", "Vibration", "Gap", "Decision")
    table = doc.add_table(rows=1, cols=len(headings))
    table.autofit = False
    widths = (0.53, 1.00, 0.50, 0.95, 0.90, 1.05, 0.60, 1.00)
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Inches(width)
    for cell, label, width in zip(table.rows[0].cells, headings, widths, strict=True):
        cell.width = Inches(width)
        cell.text = label
        _cell_fill(cell, "DCEAF5")
        _cell_border(cell)
        _cell_padding(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    for index, row in enumerate(CASES):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row[:7] + (row[7],), widths, strict=True):
            cell.width = Inches(width)
            cell.text = value
            _cell_border(cell)
            _cell_padding(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2:
                _cell_fill(cell, "F5F9FC")
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(8)
        for cell in cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build() -> Path:
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite fixture: {OUTPUT}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(11 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        if name == "Title":
            # The stock template draws a blue rule below Title in Word.
            p_pr = style._element.get_or_add_pPr()
            for border in p_pr.findall(qn("w:pBdr")):
                p_pr.remove(border)
            borders = OxmlElement("w:pBdr")
            for side_name in ("top", "left", "bottom", "right", "between"):
                side = OxmlElement(f"w:{side_name}")
                side.set(qn("w:val"), "nil")
                borders.append(side)
            p_pr.append(borders)

    doc.add_paragraph("Conveyor Signal Fault Review and Release Decisions", style="Title")
    doc.add_paragraph(
        f"{CANARY}. This fictional review record tests whether a Wiki preserves measured values, "
        "strict comparisons, source conflicts and decision reasons from a Word file. "
        "It is not operating guidance for any real machine. The decisive result is that a lane "
        "remains on hold when any required condition fails, even if the other readings appear normal."
    )

    doc.add_paragraph("Evidence and decision rules", style="Heading 1")
    doc.add_paragraph(
        "The signed incident record identifies the fault and lane. A verified snapshot provides the "
        "comparison state, while a time aligned telemetry series supplies the measurements. "
        "An operator note may explain the sequence but cannot override a contradictory measured value. "
        "If the timebase is invalid or a sample is missing, the observation window is incomplete."
    )
    doc.add_paragraph(
        "Release requires motor current strictly below 14.2 A, bearing vibration strictly below "
        "3.8 mm/s and every telemetry gap strictly below 5 s throughout a 90 s observation window. "
        "The snapshot checksum must match the signed incident record. Equality with any numeric "
        "limit fails. A measured hold remains in the audit trail even after a later correction."
    )

    doc.add_paragraph("Case measurements", style="Heading 1")
    doc.add_paragraph(
        "The table records the original six decisions. Each row concerns only its named lane. "
        "A release for lane C does not authorize release of lane B or any adjacent lane."
    )
    _table(doc)

    doc.add_paragraph("Case interpretations", style="Heading 1")
    for key, controller, lane, snapshot, current, vibration, gap, decision, reason in CASES:
        doc.add_paragraph(f"{key}  {controller} lane {lane}", style="Heading 2")
        detail = doc.add_paragraph(
            f"The signed fault record assigns controller {controller} to lane {lane} and compares it "
            f"with snapshot {snapshot}. Its 90 s trace reports current {current}, vibration "
            f"{vibration} and largest gap {gap}. The recorded decision is {decision} because {reason}. "
            "The comparison applies the strict limits independently rather than averaging the readings."
        )
        detail.paragraph_format.keep_together = True

    doc.add_paragraph("Cross document distinction", style="Heading 1")
    doc.add_paragraph(
        "D01 refers to the same fictional C17, S42 and lane B values as the companion PDF, but "
        "this Word review adds explicit equality cases D03 and D04 and evidence integrity cases "
        "D05 and D06. A knowledge graph may connect the two AMHS sources as related evidence. "
        "It should not merge their case identifiers or infer that D02 releases other lanes."
    )
    doc.add_paragraph(
        "A later note could state that D01 vibration was corrected to 3.6 mm/s. That note alone "
        "does not change the original HOLD decision. A revised signed trace and a new authorized "
        "decision would be required. This counterfactual is deliberately separate from the six "
        "recorded decisions and must not be reported as an observed release."
    )

    doc.save(OUTPUT)
    data = OUTPUT.read_bytes()
    assert data[:2] == b"PK" and len(data) > 20_000
    print({"path": str(OUTPUT), "bytes": len(data), "sha256": sha256(data).hexdigest(),
           "cases": len(CASES), "canary": CANARY})
    return OUTPUT


if __name__ == "__main__":
    build()
